import pandas as pd, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
P=pd.read_pickle("P2.pkl"); C=pd.read_pickle("C.pkl"); D=pd.read_csv("dvp.csv")
HDR=PatternFill("solid",fgColor="1F3864"); HF=Font(bold=True,color="FFFFFF"); INP=PatternFill("solid",fgColor="FFF2CC")
def hdr(ws,row=1):
    for c in ws[row]:
        if c.value is not None: c.fill=HDR; c.font=HF; c.alignment=Alignment(wrap_text=True,vertical="center")
def widths(ws,w):
    for i,x in enumerate(w,1): ws.column_dimensions[get_column_letter(i)].width=x
wb=Workbook()

# ---------- Parametros ----------
ws=wb.active; ws.title="Parametros"
ws.append(["Parámetro","Valor","Qué hace"])
ws.append(["lambda_PTS",0.7,"Cuánto del 'sube/baja' del rival en puntos se mantiene (0=ignorar defensa, 1=confiar 100%). Supuesto inicial, falta calibrar."])
ws.append(["lambda_REB",0.6,"Igual, para rebotes."])
ws.append(["lambda_AST",0.5,"Igual, para asistencias (las más ruidosas por posición)."])
ws.append(["Tope_factor",0.12,"El factor DvP nunca mueve la proyección más de ±12%."])
ws.append(["sd_PTS",0.30,"Variación típica de los puntos de un jugador = 30% de su promedio (supuesto genérico; se mide con registros por partido)."])
ws.append(["var_REB",1.3,"Rebotes: varianza = 1,3 × promedio (supuesto)."])
ws.append(["var_AST",1.4,"Asistencias: varianza = 1,4 × promedio (supuesto)."])
ws.append(["sd_PRA",0.25,"Variación típica de puntos+rebotes+asistencias = 25% de su promedio (supuesto)."])
ws.append(["EV_minimo",0.05,"Esperanza de ganancia mínima (por unidad apostada) para marcar una jugada. Alta a propósito: el modelo aún no está calibrado."])
ws.append([])
ws.append(["Fórmula del factor","","Factor = 1 + límite( continuidad_rival × lambda × (valor_rival − media_posición) / media_posición )"])
ws.append(["Ajuste de rol","","Columnas adj_* de la hoja Jugadores (calculadas fuera de la hoja): DAMP=0.30, CAP=±6%, ARR_W=0.60. Reparte la producción de los que se fueron entre los que se quedan, según su margen de minutos."])
ws.append(["Pendiente de calibrar","","lambda y los parámetros de rol son supuestos de partida. Se calibran con histórico por partido (backtest), no a ojo."])
for r in range(2,11): ws.cell(r,2).fill=INP
hdr(ws); widths(ws,[22,10,120])

# ---------- Equipos ----------
ws=wb.create_sheet("Equipos")
ws.append(["Equipo","Abrev","Continuidad","Retenido_25-26","Retenido_26-27","Principales salidas","Principales llegadas"])
for _,r in C.sort_values("Nombre").iterrows():
    ws.append([r.Nombre,r.Equipo,float(r.Continuidad),float(r.Ret_prev),float(r.Ret_curr),r.Principales_salidas,r.Principales_llegadas])
hdr(ws); widths(ws,[26,8,13,14,14,55,55]); ws.freeze_panes="A2"

# ---------- DvP ----------
ws=wb.create_sheet("DvP")
ws.append(list(D.columns))
for row in D.itertuples(index=False): ws.append(list(row))
hdr(ws); widths(ws,[26,9]+[11]*12); ws.freeze_panes="A2"

# ---------- Factor_DvP ----------
ws=wb.create_sheet("Factor_DvP")
ws.append(["Clave","Equipo rival","Pos","Continuidad rival","Dev_PTS","Dev_REB","Dev_AST","Factor_PTS","Factor_REB","Factor_AST"])
for i in range(2,152):
    ws.append([f'=B{i}&"|"&C{i}',f"=DvP!A{i}",f"=DvP!B{i}",
        f"=INDEX(Equipos!$C$2:$C$31,MATCH(B{i},Equipos!$A$2:$A$31,0))",
        f"=(DvP!C{i}-AVERAGEIFS(DvP!$C$2:$C$151,DvP!$B$2:$B$151,C{i}))/AVERAGEIFS(DvP!$C$2:$C$151,DvP!$B$2:$B$151,C{i})",
        f"=(DvP!E{i}-AVERAGEIFS(DvP!$E$2:$E$151,DvP!$B$2:$B$151,C{i}))/AVERAGEIFS(DvP!$E$2:$E$151,DvP!$B$2:$B$151,C{i})",
        f"=(DvP!D{i}-AVERAGEIFS(DvP!$D$2:$D$151,DvP!$B$2:$B$151,C{i}))/AVERAGEIFS(DvP!$D$2:$D$151,DvP!$B$2:$B$151,C{i})",
        f"=1+MAX(-Parametros!$B$5,MIN(Parametros!$B$5,D{i}*Parametros!$B$2*E{i}))",
        f"=1+MAX(-Parametros!$B$5,MIN(Parametros!$B$5,D{i}*Parametros!$B$3*F{i}))",
        f"=1+MAX(-Parametros!$B$5,MIN(Parametros!$B$5,D{i}*Parametros!$B$4*G{i}))"])
for r in ws.iter_rows(min_row=2,min_col=4,max_col=10):
    for c in r: c.number_format="0.000" if c.column>7 else "0.0%" if c.column in (5,6,7) else "0.000"
hdr(ws); widths(ws,[34,26,6,14,10,10,10,11,11,11]); ws.freeze_panes="A2"

# ---------- Jugadores ----------
ws=wb.create_sheet("Jugadores")
ws.append(["Jugador","Equipo_25-26","Equipo_26-27","Movimiento","Pos","Pos_alt","Confianza_pos","MIN","PTS","REB","AST","adj_PTS","adj_REB","adj_AST","PTS_base","REB_base","AST_base"])
Q=P.sort_values(["PTS"],ascending=False).reset_index(drop=True)
for i,r in enumerate(Q.itertuples(),start=2):
    ws.append([r.Jugador,r.Equipo_2526,r.Equipo_2627,r.Movimiento,r.Pos,r.Pos_alt,r.Confianza_pos,float(r.MIN),float(r.PTS),float(r.REB),float(r.AST),
               round(float(r.adj_PTS),3),round(float(r.adj_REB),3),round(float(r.adj_AST),3),f"=ROUND(I{i}+L{i},2)",f"=ROUND(J{i}+M{i},2)",f"=ROUND(K{i}+N{i},2)"])
NJ=len(Q)+1
hdr(ws); widths(ws,[26,11,11,24,6,7,16,7,7,7,7,9,9,9,9,9,9]); ws.freeze_panes="B2"
for r in ws.iter_rows(min_row=2,min_col=5,max_col=5):
    for c in r: c.fill=INP

# ---------- Calculadora ----------
ws=wb.create_sheet("Calculadora",0)
ws["A1"]="Calculadora de proyección de props (v0)"; ws["A1"].font=Font(bold=True,size=14)
ws["A3"]="Jugador";ws["B3"]="Luka Dončić"
ws["A4"]="Rival";ws["B4"]="Chicago Bulls"
ws["A5"]="Posición a usar (opcional)";ws["B5"]=None
for c in ("B3","B4","B5"): ws[c].fill=INP
ws["A7"]="Posición usada"; ws["B7"]=f'=IF(B5<>"",B5,INDEX(Jugadores!$E$2:$E${NJ},MATCH(B3,Jugadores!$A$2:$A${NJ},0)))'
ws["A8"]="Equipo 2026-27 del jugador"; ws["B8"]=f'=INDEX(Jugadores!$C$2:$C${NJ},MATCH(B3,Jugadores!$A$2:$A${NJ},0))'
ws["A9"]="Continuidad del rival (0-1)"; ws["B9"]="=INDEX(Equipos!$C$2:$C$31,MATCH(B4,Equipos!$A$2:$A$31,0))"; ws["B9"].number_format="0.00"
ws.append([])
ws["A11"]="Stat";ws["B11"]="Base con ajuste de rol";ws["C11"]="Factor DvP rival";ws["D11"]="Proyección";ws["E11"]="Línea de la casa";ws["F11"]="Proyección − línea";ws["G11"]="Lado"
hdr(ws,11)
cols={"PTS":("O","H"),"REB":("P","I"),"AST":("Q","J")}
for r,(st,(bc,fc)) in zip((12,13,14),cols.items()):
    ws[f"A{r}"]=st
    ws[f"B{r}"]=f"=INDEX(Jugadores!${bc}$2:${bc}${NJ},MATCH($B$3,Jugadores!$A$2:$A${NJ},0))"
    ws[f"C{r}"]=f'=INDEX(Factor_DvP!${fc}$2:${fc}$151,MATCH($B$4&"|"&$B$7,Factor_DvP!$A$2:$A$151,0))'
    ws[f"D{r}"]=f"=ROUND(B{r}*C{r},1)"
    ws[f"E{r}"].fill=INP
    ws[f"F{r}"]=f'=IF(E{r}="","",ROUND(D{r}-E{r},1))'
    ws[f"G{r}"]=f'=IF(E{r}="","",IF(F{r}>=0,"Over","Under"))'
    ws[f"B{r}"].number_format="0.0"; ws[f"C{r}"].number_format="0.000"
ws["A15"]="PRA"; ws["B15"]="=SUM(B12:B14)"; ws["D15"]="=SUM(D12:D14)"; ws["C15"]="=ROUND(D15/B15,3)"; ws["E15"].fill=INP
ws["F15"]='=IF(E15="","",ROUND(D15-E15,1))'; ws["G15"]='=IF(E15="","",IF(F15>=0,"Over","Under"))'
ws["B15"].number_format="0.0"
ws["A17"]="OJO: la diferencia con la línea NO es una ventaja por sí sola. Falta compararla con la variabilidad real del jugador y con la cuota."
ws["A18"]="Celdas amarillas = las que puedes cambiar. Los parámetros están en la hoja Parametros."
ws["A17"].font=Font(italic=True,color="C00000")
widths(ws,[30,24,16,13,16,18,8])
dv1=DataValidation(type="list",formula1=f"=Jugadores!$A$2:$A${NJ}",allow_blank=False); ws.add_data_validation(dv1); dv1.add("B3")
dv2=DataValidation(type="list",formula1="=Equipos!$A$2:$A$31",allow_blank=False); ws.add_data_validation(dv2); dv2.add("B4")
dv3=DataValidation(type="list",formula1='"PG,SG,SF,PF,C"',allow_blank=True); ws.add_data_validation(dv3); dv3.add("B5")

# ---------- Bloque de probabilidad en la Calculadora ----------
ws=wb["Calculadora"]
ws["A20"]="PROBABILIDAD DEL MODELO vs PROBABILIDAD DE LA CASA"; ws["A20"].font=Font(bold=True,size=12)
ws["A21"]="Escribe la línea arriba (columna E) y aquí las cuotas DECIMALES del over y del under. Las filas de ejemplo (asistencias) son solo una demostración: bórralas o cámbialas."
ws["A21"].font=Font(italic=True)
heads=["Stat","Proyección","Línea","Cuota over","Cuota under","P modelo Over","P modelo Under","Implícita over (con margen)","Implícita over (sin margen)","Ventaja vs justa","EV over (por unidad)","EV under (por unidad)","Lectura"]
for j,h in enumerate(heads,1): ws.cell(22,j,h)
hdr(ws,22); ws.row_dimensions[22].height=45
src={23:12,24:13,25:14,26:15}   # filas de la tabla de arriba
names={23:"PTS",24:"REB",25:"AST",26:"PRA"}
for r,sr in src.items():
    ws[f"A{r}"]=names[r]; ws[f"B{r}"]=f"=D{sr}"; ws[f"C{r}"]=f'=IF(E{sr}="","",E{sr})'
    ws[f"D{r}"].fill=INP; ws[f"E{r}"].fill=INP
    if names[r] in ("PTS","PRA"):
        k="$B$6" if names[r]=="PTS" else "$B$9"
        ws[f"F{r}"]=f'=IF(OR($C{r}="",$B{r}=""),"",1-NORMDIST(INT($C{r})+0.5,$B{r},Parametros!{k}*$B{r},TRUE))'
        ws[f"G{r}"]=f'=IF(OR($C{r}="",$B{r}=""),"",NORMDIST(-INT(-$C{r})-0.5,$B{r},Parametros!{k}*$B{r},TRUE))'
    else:
        v="$B$7" if names[r]=="REB" else "$B$8"
        ws[f"F{r}"]=f'=IF(OR($C{r}="",$B{r}=""),"",1-BETADIST(1/Parametros!{v},$B{r}/(Parametros!{v}-1),INT($C{r})+1))'
        ws[f"G{r}"]=f'=IF(OR($C{r}="",$B{r}=""),"",BETADIST(1/Parametros!{v},$B{r}/(Parametros!{v}-1),-INT(-$C{r})))'
    ws[f"H{r}"]=f'=IF($D{r}="","",1/$D{r})'
    ws[f"I{r}"]=f'=IF(OR($D{r}="",$E{r}=""),"",(1/$D{r})/((1/$D{r})+(1/$E{r})))'
    ws[f"J{r}"]=f'=IF(OR($F{r}="",$I{r}=""),"",$F{r}-$I{r})'
    ws[f"K{r}"]=f'=IF(OR($F{r}="",$D{r}="",$G{r}=""),"",$F{r}*($D{r}-1)-$G{r})'
    ws[f"L{r}"]=f'=IF(OR($G{r}="",$E{r}="",$F{r}=""),"",$G{r}*($E{r}-1)-$F{r})'
    ws[f"M{r}"]=f'=IF(OR($K{r}="",$L{r}=""),"",IF(MAX($K{r},$L{r})>=Parametros!$B$10,IF($K{r}>=$L{r},"Over","Under"),"Sin ventaja suficiente"))'
    for c in "FGHIJKL": ws[f"{c}{r}"].number_format="0.0%"
    ws[f"B{r}"].number_format="0.0"; ws[f"D{r}"].number_format="0.00"; ws[f"E{r}"].number_format="0.00"
# ejemplo de demostración (asistencias)
ws["E14"]=8.5; ws["D25"]=1.87; ws["E25"]=1.95
ws["A28"]="Cómo leerlo"; ws["A28"].font=Font(bold=True)
notes=["P modelo = probabilidad (según el modelo) de que el jugador supere / se quede debajo de la línea. En línea entera (ej. 9) la diferencia hasta 100% es 'push' (te devuelven la apuesta).",
"Implícita con margen = 1 / cuota. Es el punto de equilibrio: necesitas acertar MÁS que eso para ganar dinero a largo plazo.",
"Implícita sin margen = lo que la casa 'cree' de verdad, quitando su comisión. Sirve para medir la ventaja (columna J).",
"EV = ganancia esperada por cada unidad apostada = P × (cuota − 1) − (1 − P). Si es negativa, esa apuesta pierde dinero a largo plazo aunque la ventaja de la columna J sea positiva.",
"El modelo NO está calibrado todavía: usa la hoja Registro para anotar lo que pasa y comprobar si sus probabilidades se cumplen."]
for i,t in enumerate(notes): ws[f"A{29+i}"]="• "+t
ws["A35"]="Conversor de cuota americana → decimal"; ws["A35"].font=Font(bold=True)
ws["A36"]="Cuota americana (ej. -115 o +120)"; ws["B36"]=-115; ws["B36"].fill=INP
ws["A37"]="Cuota decimal"; ws["B37"]='=IF(B36="","",IF(B36>0,1+B36/100,1+100/ABS(B36)))'; ws["B37"].number_format="0.000"
for c,w in zip("HIJKLM",[16,16,13,14,14,24]): ws.column_dimensions[c].width=w

# ---------- Registro de jugadas (para calibrar) ----------
wr=wb.create_sheet("Registro")
wr["A1"]="Registro de jugadas (para medir si el modelo acierta lo que dice)"; wr["A1"].font=Font(bold=True,size=13)
wr["A2"]="Anota cada jugada que analices (aunque no apuestes). Pega la 'P modelo' de ese lado, y al terminar el partido escribe el resultado real."
wr["A4"]="Jugadas resueltas";wr["B4"]="=COUNT(I9:I508)"
wr["A5"]="% de acierto real";wr["B5"]='=IF(B4=0,"",AVERAGE(I9:I508))'
wr["A6"]="Prob. media que daba el modelo";wr["B6"]='=IF(B4=0,"",SUMPRODUCT(ISNUMBER(I9:I508)*G9:G508)/B4)'
wr["D4"]="Error de Brier (menor = mejor; 0,25 = igual que lanzar una moneda)";wr["F4"]='=IF(B4=0,"",AVERAGE(J9:J508))'
wr["D5"]="Si % de acierto real ≈ prob. media del modelo, está bien calibrado.";
for c in ("B5","B6"): wr[c].number_format="0.0%"
wr["F4"].number_format="0.000"
for j,h in enumerate(["Fecha","Jugador","Stat","Lado (Over/Under)","Línea","Cuota decimal","P modelo (de ese lado)","Resultado real","Acertó (1/0)","Brier"],1): wr.cell(8,j,h)
hdr(wr,8)
for r in range(9,509):
    wr[f"I{r}"]=f'=IF(OR(H{r}="",E{r}="",D{r}=""),"",IF(H{r}=E{r},"push",IF(D{r}="Over",IF(H{r}>E{r},1,0),IF(H{r}<E{r},1,0))))'
    wr[f"J{r}"]=f'=IF(AND(ISNUMBER(I{r}),G{r}<>""),(G{r}-I{r})^2,"")'
    wr[f"G{r}"].number_format="0.0%"
dvl=DataValidation(type="list",formula1='"Over,Under"',allow_blank=True); wr.add_data_validation(dvl); dvl.add("D9:D508")
dvs=DataValidation(type="list",formula1='"PTS,REB,AST,PRA"',allow_blank=True); wr.add_data_validation(dvs); dvs.add("C9:C508")
for c,w in zip("ABCDEFGHIJ",[32,24,8,16,8,12,22,14,12,9]): wr.column_dimensions[c].width=w
wr.freeze_panes="A9"

wb.save("cruce_raw.xlsx"); print("ok",NJ-1,"jugadores")
